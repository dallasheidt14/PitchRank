import csv
from dataclasses import replace
from io import BytesIO, StringIO

from openpyxl import load_workbook

from src.tournaments.seeding_content import build_director_cohort, director_csv
from src.tournaments.seeding_result_explanation import explain_results
from src.tournaments.seeding_sheet import CohortSheet, SheetTeam, render_sheet_html
from src.tournaments.seeding_tiers import TierEntrant, build_cheat_sheet_analysis
from src.tournaments.seeding_workbook import build_seeding_workbook


def test_game_examples_include_contrary_results_and_retain_sparse_warning():
    profile = {"games": 2, "distinctOpponents": 2, "latestGameDate": "2026-10-01", "limited": True,
               "componentEntrants": 1, "influentialGames": [
                   {"margin": 2, "residual": 1, "opponent_name": "Alpha", "opponent_club": "Club",
                    "date": "2026-10-01", "fixture": {"goals": 4}},
                   {"margin": -2, "residual": -1, "opponent_name": "Beta", "date": "2026-09-20",
                    "fixture": {"goals": 4}}]}
    reason = explain_results(profile)
    assert "3-1 vs Club Alpha (2026-10-01)" in reason
    assert "1-3 vs Beta (2026-09-20)" in reason
    assert "Limited evidence" in reason and "No shared-result link" in reason
    assert "saved rating" in explain_results({"games": 0})


def test_every_director_export_preserves_reason_compromise_and_coverage():
    analysis = build_cheat_sheet_analysis([TierEntrant("a", "=Alpha", .8)], {})
    reason = "3-1 vs Club (2026-10-01). Limited evidence."
    analysis = replace(analysis, ordering_evidence={"teams": {"a": {"director_reason": reason,
        "cautions": ["8 games lack usable opponent ratings", "Field sensitivity: seeds 1–4"]}}},
        tier_guidance={"status": "review_required_compromise", "cuts": [1],
            "divisions": [{"label": "Tier 1", "start_seed": 1, "end_seed": 1, "size": 1,
                           "pool_sizes": [1], "minimum_games": 0}], "warnings": ["Review worst matchup."]})
    sheet = CohortSheet("u13", "Female", (SheetTeam("=Alpha", "Club", .8, entrant_id="a", requested_flight="Gold", listed_division="League A", record="8-3-1"),),
                        (SheetTeam("Unplaced", "Club", entrant_id="b"),), tier_analysis=analysis)
    content = build_director_cohort(sheet)
    assert any("8 games lack usable opponent ratings" in note and "Field sensitivity: seeds 1–4" in note
               for note in content.notes)
    assert "Provisional tier break" in content.rows[0].break_label
    assert content.rows[0].recommendation_reason == reason
    rows = list(csv.DictReader(StringIO(director_csv(sheet).decode("utf-8-sig"))))
    assert len(rows) == 2 and rows[0]["Team"] == "'=Alpha"
    assert rows[0]["Coverage"] == "1 of 2 teams assessed"
    assert "Review-required compromise" in rows[0]["Recommendation"]
    assert "Requested: Gold" in rows[0]["Roster context"] and "Listed: League A" in rows[0]["Roster context"]
    assert rows[1]["Seed"] == ""
    assert "8 games lack usable opponent ratings" in rows[0]["Warnings and formats"]
    html = render_sheet_html("Cup", [sheet], generated_on="2026-10-18", ranking_run="2026-10-17")
    assert reason not in html and "Unplaced" in html and "Matchup review needed" in html
    assert "What to know" not in html
    assert "Field sensitivity" not in html
    assert "8 games lack usable opponent ratings" in html
    assert "Requested: Gold" in html and "Listed: League A" in html
    assert '<th class="num">PowerScore</th>' in html
    assert "8-3-1" in html and "Record<br>(W-L-D)" in html
    assert rows[0]["Record (W-L-D)"] == "8-3-1"
    assert rows[0]["PowerScore"] == "80.0"
    book = load_workbook(BytesIO(build_seeding_workbook("Cup", [sheet], generated_on="2026-10-18",
                                                     ranking_run="2026-10-17")))
    assert any("8 games lack usable opponent ratings" in str(cell.value)
               for row in book.active for cell in row)
    assert "Requested: Gold" in book.active["B7"].value and "Listed: League A" in book.active["B7"].value
    assert reason not in book.active["G7"].value
    assert "Tier 1" in book.active["G7"].value
    assert not any(book.active.column_dimensions[col].hidden for col in ("D", "E", "F", "L"))
    assert book.active["L7"].value == "8-3-1"
    assert book.active["B8"].value == "Unplaced"


def test_legacy_embedded_format_keeps_unknown_game_counts_exportable():
    from pathlib import Path

    from src.tournaments.seeding_format_library import load_format_library
    from src.tournaments.seeding_format_preferences import format_summary

    library = load_format_library(Path(__file__).resolve().parents[2] / "config/matchbalance_format_library_v1.json")
    template = library.templates[0]
    summary = format_summary(template)
    assert summary["maximum_pair_meetings"] is None
    analysis = build_cheat_sheet_analysis([TierEntrant("a", "Alpha", .8)], {})
    analysis = replace(analysis, tier_guidance={"divisions": [
        {"label": "Tier 1", "size": 1, "start_seed": 1, "end_seed": 1,
         "pool_sizes": list(template.pool_sizes), "minimum_games": summary["minimum_games"],
         "selected_format": {**summary, "maximum_games": None}}]})
    sheet = CohortSheet("u13", "Female", (SheetTeam("Alpha", "Club", .8, entrant_id="a"),), (), analysis)
    content = build_director_cohort(sheet)
    assert "maximum not specified" in content.format_notes[0]
    assert "includes playoffs" in content.format_notes[0]
    html = render_sheet_html("Legacy Cup", [sheet], generated_on="2026-10-18", ranking_run="2026-10-17")
    assert "maximum not specified" not in html and "up to None" not in html
    workbook = load_workbook(BytesIO(build_seeding_workbook("Legacy Cup", [sheet],
        generated_on="2026-10-18", ranking_run="2026-10-17")))
    assert workbook.active["B7"].value == "Alpha"


def test_manually_seeded_review_team_keeps_identity_warning_in_director_view():
    analysis = build_cheat_sheet_analysis(
        [TierEntrant("a", "Review team", review_reason="Confirm identity")], {}, manual_order=("a",))
    sheet = CohortSheet("u13", "Female", (),
        (SheetTeam("Review team", "Club", entrant_id="a", review_reason="Confirm identity"),), analysis)
    content = build_director_cohort(sheet)
    assert content.rows[0].seed == 1
    assert "Review team: Confirm identity." in content.director_notes
    html = render_sheet_html("Cup", [sheet], generated_on="2026-10-18", ranking_run="2026-10-17")
    assert "Review team: Confirm identity." in html
    workbook = load_workbook(BytesIO(build_seeding_workbook("Cup", [sheet],
        generated_on="2026-10-18", ranking_run="2026-10-17")))
    assert any("Review team: Confirm identity." == cell.value for row in workbook.active for cell in row)


def test_tier_stability_warning_survives_unchanged_seed_positions():
    analysis = build_cheat_sheet_analysis([TierEntrant("a", "Alpha", .8)], {})
    for field in ("changed_tier_memberships", "cluster_changed_tier_memberships"):
        evidence = {"teams": {"a": {"director_reason": "Game evidence"}},
                    "game_sensitivity": {field: 1, "seed_ranges": {"a": [1, 1]}}}
        sheet = CohortSheet("u13", "Female", (SheetTeam("Alpha", "Club", .8, entrant_id="a"),), (),
                            replace(analysis, ordering_evidence=evidence))
        content = build_director_cohort(sheet)
        assert any("Some placements depend on a small number of results" in note for note in content.director_notes)
        assert all("Several positions change" not in note for note in content.director_notes)


def test_director_review_list_preserves_sparse_teams_without_operator_format_details():
    analysis = build_cheat_sheet_analysis([TierEntrant("a", "Review & Co", .8)], {})
    analysis = replace(analysis, limited_history=("a",), tier_guidance={
        "status": "passing", "cuts": [], "warnings": ["Confirm missing opponent results."],
        "divisions": [{"label": "Tier 1", "start_seed": 1, "end_seed": 1, "size": 1,
                       "pool_sizes": [1], "minimum_games": 0}],
    })
    sheet = CohortSheet("u13", "Male", (SheetTeam("Review & Co", "Club", .8, entrant_id="a"),), (), analysis)
    content = build_director_cohort(sheet)
    assert content.review_teams == ("Review & Co",)
    assert content.format_notes
    args = dict(generated_on="2026-10-20", ranking_run="2026-10-20")
    document = render_sheet_html("Cup", [sheet], **args)
    checklist = document.split('<aside class="guidance">', 1)[1]
    assert '<ul class="review-team-list"><li>Review &amp; Co</li></ul>' in checklist
    assert "We have too few usable recent results" in checklist
    assert "take a closer look" in checklist
    assert "Confirm missing opponent results." in checklist
    assert "Tier 1 —" not in checklist
    tab = load_workbook(BytesIO(build_seeding_workbook("Cup", [sheet], **args))).active
    assert any(cell.value == "Review & Co" for row in tab.iter_rows(min_row=9) for cell in row)
    assert any(cell.value == "Confirm missing opponent results." for row in tab for cell in row)
    assert not any(str(cell.value).startswith("Tier 1 —") for row in tab for cell in row)
    csv_rows = list(csv.DictReader(StringIO(director_csv(sheet).decode("utf-8-sig"))))
    assert "Teams to review: Review & Co" in csv_rows[0]["Warnings and formats"]


def test_director_exports_translate_operations_without_changing_saved_reason():
    from src.tournaments.seeding_tiers import NO_CURRENT_RATING

    raw = ("Rankings do not currently include usable game history for this team. "
           "Refresh results, then run rankings and rebuild.")
    entrants = [TierEntrant("a", "Missing history", review_reason=raw, review_status=NO_CURRENT_RATING)]
    analysis = build_cheat_sheet_analysis(entrants, {})
    team = SheetTeam("Missing history", "Club", entrant_id="a", review_reason=raw,
                     review_kind="no_current_rating")
    expected = ("Not enough data to recommend a placement. "
                "Please check recent results or ask the coach about the team's level.")
    args = dict(generated_on="2026-10-20", ranking_run="2026-10-20")
    for manual in (False, True):
        current = build_cheat_sheet_analysis(entrants, {}, manual_order=("a",)) if manual else analysis
        sheet = CohortSheet("u13", "Male", (), (team,), current)
        row = build_director_cohort(sheet).rows[0]
        assert row.team.review_reason == raw
        assert row.recommendation_reason == raw or raw in row.recommendation_reason
        assert row.placement_explanation == expected
        document = render_sheet_html("Cup", [sheet], **args)
        tab = load_workbook(BytesIO(build_seeding_workbook("Cup", [sheet], **args))).active
        csv_text = director_csv(sheet).decode("utf-8-sig")
        for output in (document, str([cell.value for r in tab for cell in r]), csv_text):
            assert "Not enough data to recommend a placement." in output
            assert "run rankings" not in output and "Refresh results" not in output
        assert team.review_reason == raw
