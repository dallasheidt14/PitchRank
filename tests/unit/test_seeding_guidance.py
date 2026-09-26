"""Director claims must be narrower than internal matchup-limit diagnostics."""

from dataclasses import replace
from io import BytesIO
from itertools import combinations

import pytest
from openpyxl import load_workbook

from src.tournaments.compare_predictor_bridge import ComparePrediction
from src.tournaments.seeding_content import build_director_cohort, export_fingerprint
from src.tournaments.seeding_sheet import CohortSheet, SheetTeam, render_sheet_html
from src.tournaments.seeding_tiers import TierEntrant, TierPolicy, build_cheat_sheet_analysis
from src.tournaments.seeding_workbook import build_seeding_workbook


def prediction(margin, risk=.1):
    return ComparePrediction(
        "team_a", .6, .3, .1, {"teamA": 2, "teamB": 1}, margin, max(abs(margin), .5), risk, "low", .4,
    )


def cohort(count, *, gap=None, bottom=False):
    entrants = [TierEntrant(str(i), f"Team {i}", .9 - i * .01) for i in range(count)]
    if gap:
        boundary = count - gap if bottom else gap
        entrants = [replace(e, power_score=e.power_score - (.3 if i >= boundary else 0))
                    for i, e in enumerate(entrants)]
    pairs = {(a.entrant_id, b.entrant_id): prediction(
        3 if gap and int(a.entrant_id) < boundary <= int(b.entrant_id) else .1,
        .5 if gap and int(a.entrant_id) < boundary <= int(b.entrant_id) else .1,
    ) for a, b in combinations(entrants, 2)}
    analysis = build_cheat_sheet_analysis(entrants, pairs)
    sheet = CohortSheet("u14", "Male", tuple(
        SheetTeam(e.team_name, "Club", e.power_score, entrant_id=e.entrant_id) for e in entrants
    ), (), analysis)
    return sheet, entrants, pairs


def consensus_case(compare_strengths, *, count=7, limited=()):
    ids = [str(index) for index in range(count)]
    entrants = [
        TierEntrant(
            entrant_id,
            f"Team {entrant_id}",
            .90 - index * .01,
            limited_history=entrant_id in limited,
            evidence_game_count=12 + index,
        )
        for index, entrant_id in enumerate(ids)
    ]
    pairs = {
        (first, second): prediction(compare_strengths[first] - compare_strengths[second])
        for first, second in combinations(ids, 2)
    }
    return entrants, pairs


@pytest.mark.parametrize("count", [0, 1, 2, *range(3, 17), 32])
def test_counts_preserved_without_equivalence_or_format_claims(count):
    sheet, _, _ = cohort(count)
    content = build_director_cohort(sheet)
    assert len(content.rows) == count
    assert [row.seed for row in content.rows] == list(range(1, count + 1))
    assert len({row.team.entrant_id for row in content.rows}) == count
    assert content.notes == ()
    assert all(row.observation == "" for row in content.rows)
    assert all(sheet.tier_analysis.marker_for_seed(seed) == "" for seed in range(1, count + 1))
    document = render_sheet_html("Event", [sheet], generated_on="2026-09-20", ranking_run="2026-09-19")
    assert "Close range" not in document and "balanced across pools" not in document
    assert "singleton" not in document and "one team;" not in document
    assert "advancement" not in document and "games guaranteed" not in document
    workbook = load_workbook(BytesIO(build_seeding_workbook(
        "Event", [sheet], generated_on="2026-09-20", ranking_run="2026-09-19",
    )))
    tab = workbook.active
    for index, row in enumerate(content.rows, 7):
        assert tab.cell(index, 1).value == row.seed
        assert tab.cell(index, 2).value == row.team.team_name
        assert tab.cell(index, 4).value == pytest.approx(row.score)
        assert tab.cell(index, 6).value is None
    assert "Director notes" not in [tab.cell(i, 1).value for i in range(7, tab.max_row + 1)]


def test_overlapping_ranges_never_label_seed_one_and_thirteen_as_equivalent():
    sheet, _, _ = cohort(15)
    ranges = sheet.tier_analysis.close_ranges
    assert any(a.start_seed < b.start_seed <= a.end_seed < b.end_seed
               for a, b in combinations(ranges, 2))
    content = build_director_cohort(sheet)
    assert content.rows[0].seed == 1 and content.rows[12].seed == 13
    assert content.rows[0].score > content.rows[12].score
    assert content.rows[0].observation == content.rows[12].observation == ""
    assert content.notes == ()


def test_close_neighbors_do_not_hide_unsafe_endpoints_in_internal_evidence():
    _, entrants, pairs = cohort(3)
    pairs[("0", "2")] = prediction(2.5)
    analysis = build_cheat_sheet_analysis(entrants, pairs)
    assert [(r.start_seed, r.end_seed) for r in analysis.close_ranges] == [(1, 2), (2, 3)]


@pytest.mark.parametrize("gap,bottom,expected", [
    (1, False, "Seed 1 is 31.0 points above seed 2."),
    (2, False, "Seed 2 is 31.0 points above seed 3."),
    (1, True, "Seed 7 is 31.0 points above seed 8."),
    (2, True, "Seed 6 is 31.0 points above seed 7."),
])
def test_standouts_name_the_entire_end_segment(gap, bottom, expected):
    sheet, _, _ = cohort(8, gap=gap, bottom=bottom)
    analysis = sheet.tier_analysis
    assert analysis.notes == (expected,)
    assert analysis.breaks == ()
    assert len(analysis.standouts) == 1
    standout = analysis.standouts[0]
    assert standout.standout_end_seed - standout.standout_start_seed + 1 == gap


def test_all_window_sizes_must_agree_and_rejected_evidence_survives():
    _, entrants, pairs = cohort(8, gap=4)
    pairs[("3", "4")] = prediction(-.1)
    analysis = build_cheat_sheet_analysis(entrants, pairs)
    evidence = {w.size: w for w in analysis.boundary_windows if w.after_seed == 4}
    assert set(evidence) == {3, 4, 5}
    assert not evidence[3].supported
    assert evidence[4].supported and evidence[5].supported
    assert evidence[3].upper_ids == ("3",) and evidence[3].lower_ids == ("4", "5")
    assert evidence[3].favored_fraction == .5
    assert 4 not in [b.after_seed for b in analysis.breaks]


def test_clear_gap_has_one_specific_observation_across_pdf_and_excel():
    sheet, _, _ = cohort(8, gap=4)
    content = build_director_cohort(sheet, "=Keep this literal.")
    assert content.rows[3].observation == "Score step: 31.0 points between seeds 4 and 5."
    assert content.rows[3].strength_break_after
    assert not content.rows[4].strength_break_after
    assert content.notes == ("Score step: 31.0 points between seeds 4 and 5.", "=Keep this literal.")
    args = dict(generated_on="2026-09-20", ranking_run="2026-09-19",
                operator_notes={("u14", "Male"): "=Keep this literal."})
    document = render_sheet_html("=Event", [sheet], **args)
    tab = load_workbook(BytesIO(build_seeding_workbook("=Event", [sheet], **args))).active
    assert 'data-entrant="3" class="strength-break"' in document
    assert tab.cell(10, 6).value == content.rows[3].observation
    assert tab.cell(10, 1).border.bottom.style == "medium"
    assert tab["A1"].data_type == "s"
    for note in content.notes:
        cell = next(cell for row in tab for cell in row if cell.value == note)
        assert cell.data_type == "s" and note in document


def test_fingerprint_includes_notes_policy_and_content_not_only_prediction_date():
    sheet, _, _ = cohort(6)
    args = dict(generated_on="2026-09-20", ranking_run="2026-09-19", operator_notes={},
                analysis_version=2, policy={"max_expected_margin": 2.0})
    original = export_fingerprint("Event", [sheet], **args)
    assert export_fingerprint("Event", [sheet], **args) == original
    for changed in [dict(operator_notes={("u14", "Male"): "New note"}),
                    dict(policy={"max_expected_margin": 1.5}), dict(analysis_version=3),
                    dict(ranking_run="2026-09-20")]:
        assert export_fingerprint("Event", [sheet], **{**args, **changed}) != original
    assert export_fingerprint("Event · DRAFT", [sheet], **args) != original


def test_many_supported_gaps_show_at_most_three_observations():
    _, entrants, _ = cohort(24)
    entrants = [replace(e, power_score=.95 - (i // 4) * .12 - (i % 4) * .001)
                for i, e in enumerate(entrants)]
    pairs = {(a.entrant_id, b.entrant_id): prediction(
        3 if int(a.entrant_id) // 4 < int(b.entrant_id) // 4 else .1,
        .5 if int(a.entrant_id) // 4 < int(b.entrant_id) // 4 else .1,
    ) for a, b in combinations(entrants, 2)}
    analysis = build_cheat_sheet_analysis(entrants, pairs)
    assert len(analysis.breaks) == len(analysis.notes) == 3
    assert len({window.after_seed for window in analysis.boundary_windows if window.supported}) > 3


def test_workbook_keeps_secondary_identity_and_roster_context_on_team_row():
    team = SheetTeam("=Registered name", "Club", .7, entrant_id="0",
                     pitchrank_team_name="=Database name", plays_up=True,
                     play_up_from_age_group="u13", requested_flight="Gold", listed_division="Boys U14")
    sheet = CohortSheet("u14", "Male", (team,), ())
    content = build_director_cohort(sheet)
    args = dict(generated_on="2026-09-20", ranking_run="2026-09-19")
    document = render_sheet_html("Event", [sheet], **args)
    tab = load_workbook(BytesIO(build_seeding_workbook("Event", [sheet], **args))).active
    assert tab["B7"].value.splitlines() == list(content.rows[0].name_lines)
    assert tab["B7"].data_type == "s"
    assert tab["D7"].value == 70
    assert all(text in document for text in content.rows[0].name_lines)
    assert tab.row_dimensions[7].height >= 75
    assert tab.auto_filter.ref is None
    assert [table.autoFilter.ref for table in tab.tables.values()] == ["A6:K7"]


def test_score_step_does_not_require_a_likely_blowout():
    _, entrants, pairs = cohort(8)
    entrants = [replace(e, power_score=e.power_score - (.04 if i >= 4 else 0))
                for i, e in enumerate(entrants)]
    pairs = {key: prediction(.4, .08) for key in pairs}
    analysis = build_cheat_sheet_analysis(entrants, pairs)
    assert [item.after_seed for item in analysis.breaks] == [4]
    assert analysis.breaks[0].score_gap == pytest.approx(.05)
    assert all(w.over_limit_fraction == 0 for w in analysis.boundary_windows)
    assert analysis.ordered_ids == tuple(str(i) for i in range(8))


@pytest.mark.parametrize("equal", [True, False])
def test_predictions_alone_do_not_create_score_steps(equal):
    _, entrants, pairs = cohort(8)
    if equal:
        entrants = [replace(e, power_score=.7) for e in entrants]
    analysis = build_cheat_sheet_analysis(entrants, {key: prediction(4, .6) for key in pairs})
    assert not analysis.breaks and not analysis.standouts
    assert not analysis.notes


@pytest.mark.parametrize("scores,expected", [
    ([.90, .50, .49], "Seed 1 is 40.0 points above seed 2."),
    ([.90, .89, .20], "Seed 2 is 69.0 points above seed 3."),
    ([.90, .70, .50], None),
    ([.70, .70, .70], None),
    ([.70, .695, .69], None),
])
def test_three_team_cohorts_distinguish_standouts_from_gradual_scores(scores, expected):
    entrants = [TierEntrant(str(i), f"Team {i}", score) for i, score in enumerate(scores)]
    pairs = {(a.entrant_id, b.entrant_id): prediction(.5) for a, b in combinations(entrants, 2)}
    analysis = build_cheat_sheet_analysis(entrants, pairs)
    assert analysis.notes == ((expected,) if expected else ())
    assert len(analysis.standouts) == bool(expected)
    assert not analysis.breaks


def test_only_material_reversals_need_operator_review_without_reordering():
    _, entrants, pairs = cohort(5)
    pairs[("0", "4")] = prediction(-1.25)
    pairs[("1", "2")] = prediction(-.2)
    entrants[1] = replace(entrants[1], limited_history=True)
    analysis = build_cheat_sheet_analysis(entrants, pairs)
    assert analysis.ordered_ids == ("0", "1", "2", "3", "4")
    assert analysis.limited_history == ("1",)
    assert [
        (c.upper_seed, c.lower_seed, c.lower_expected_advantage) for c in analysis.placement_checks
    ] == [(1, 5, 1.25)]
    assert not any("favored" in note for note in analysis.notes)


def test_material_reversal_threshold_is_independent_of_competitive_limit():
    _, entrants, pairs = cohort(5)
    pairs[("1", "2")] = prediction(-1.1)
    policy = TierPolicy(
        max_expected_margin=2.5,
        material_reversal_expected_goal_difference=1.0,
    )

    analysis = build_cheat_sheet_analysis(entrants, pairs, policy)

    assert [
        (check.upper_seed, check.lower_seed, check.lower_expected_advantage)
        for check in analysis.placement_checks
    ] == [(2, 3, 1.1)]


def test_local_consensus_proposes_review_only_when_shared_neighborhood_is_stable():
    entrants, pairs = consensus_case({
        "0": 11, "1": 8, "2": 10, "3": 7, "4": 6, "5": 5, "6": 4,
    })

    analysis = build_cheat_sheet_analysis(entrants, pairs)

    assert analysis.ordered_ids == ("0", "1", "2", "3", "4", "5", "6")
    proposal = next(item for item in analysis.local_consensus_checks if item.entrant_id == "2")
    assert proposal.supported and proposal.stable
    assert (proposal.baseline_seed, proposal.proposed_seed, proposal.compared_with_seed) == (3, 2, 2)
    assert proposal.direct_expected_advantage == pytest.approx(2)
    assert proposal.support_fraction == 1
    assert proposal.average_profile_advantage == pytest.approx(2)
    assert proposal.tested_window_sizes == (5, 6, 7)
    assert proposal.minimum_game_count == 12
    assert proposal.evidence_quality == "established"


def test_isolated_head_to_head_reversal_does_not_earn_a_move_proposal():
    entrants, pairs = consensus_case({
        "0": 11, "1": 10, "2": 9, "3": 8, "4": 7, "5": 6, "6": 5,
    })
    pairs[("1", "2")] = prediction(-1.25)

    analysis = build_cheat_sheet_analysis(entrants, pairs)

    review = next(item for item in analysis.local_consensus_checks if item.entrant_id == "2")
    assert not review.supported
    assert not review.stable
    assert any("shared-neighborhood" in blocker for blocker in review.blockers)
    assert analysis.ordered_ids == ("0", "1", "2", "3", "4", "5", "6")


def test_limited_history_blocks_a_consensus_move_without_treating_the_team_as_weak():
    entrants, pairs = consensus_case(
        {"0": 11, "1": 8, "2": 10, "3": 7, "4": 6, "5": 5, "6": 4}, limited={"2"},
    )

    analysis = build_cheat_sheet_analysis(entrants, pairs)

    review = next(item for item in analysis.local_consensus_checks if item.entrant_id == "2")
    assert not review.supported
    assert review.evidence_quality == "limited"
    assert any("limited ranked history" in blocker for blocker in review.blockers)
    assert analysis.ordered_ids[2] == "2"


def test_consensus_movement_cap_is_anchored_to_the_original_powerscore_order():
    entrants, pairs = consensus_case({
        "0": 9, "1": 8, "2": 7, "3": 12, "4": 6, "5": 5, "6": 4,
    })

    analysis = build_cheat_sheet_analysis(entrants, pairs)

    distant = next(
        item for item in analysis.local_consensus_checks
        if item.entrant_id == "3" and item.compared_with_id == "0"
    )
    assert not distant.supported
    assert distant.proposed_seed is None
    assert any("2-seed pilot cap" in blocker for blocker in distant.blockers)
    assert analysis.ordered_ids == ("0", "1", "2", "3", "4", "5", "6")


def test_one_influential_neighbor_cannot_create_a_stable_consensus():
    entrants, pairs = consensus_case({str(index): 0 for index in range(5)}, count=5)
    pairs.update({
        ("0", "1"): prediction(5),
        ("0", "2"): prediction(2),
        ("1", "2"): prediction(-1.25),
        ("1", "3"): prediction(0),
        ("2", "3"): prediction(1),
        ("1", "4"): prediction(.5),
        ("2", "4"): prediction(0),
    })

    analysis = build_cheat_sheet_analysis(entrants, pairs)

    review = next(item for item in analysis.local_consensus_checks if item.entrant_id == "2")
    assert review.support_fraction == pytest.approx(2 / 3)
    assert review.average_profile_advantage > 0
    assert not review.stable and not review.supported
    assert any("unusually influential" in blocker for blocker in review.blockers)


def test_limited_history_keeps_seed_and_literal_row_context_with_fixed_score_scale():
    sheet, _, _ = cohort(8, gap=4)
    teams = list(sheet.rated)
    teams[0] = replace(teams[0], status="Not Enough Ranked Games")
    sheet = replace(sheet, rated=tuple(teams))
    content = build_director_cohort(sheet)
    assert content.rows[0].seed == 1
    assert content.rows[0].display_status == "Seeded · Limited history"
    args = dict(generated_on="2026-09-20", ranking_run="2026-09-19")
    document = render_sheet_html("Event", [sheet], **args)
    tab = load_workbook(BytesIO(build_seeding_workbook("Event", [sheet], **args))).active
    assert 'style="width:90.00%"' in document
    assert 'class="flag">Limited history' in document
    assert tab["D7"].value == 90 and tab["G7"].value == content.rows[0].display_status
    assert "Seeded: 8" in tab["A5"].value
    rule = tab.conditional_formatting[next(iter(tab.conditional_formatting))][0]
    assert [(x.type, x.val) for x in rule.dataBar.cfvo] == [("num", 0), ("num", 100)]
    assert tab["A7"].border.bottom.style == "hair"
    # A director sorting the whole table keeps seed, status and marker together.
    table_rows = sorted([tuple(cell.value for cell in row) for row in tab.iter_rows(min_row=7, max_row=14)],
                        key=lambda row: row[1], reverse=True)
    marked = next(row for row in table_rows if row[0] == 4)
    assert marked[5] == content.rows[3].observation
