from io import BytesIO
from itertools import combinations

from openpyxl import load_workbook
from openpyxl.worksheet.filters import AutoFilter
import pytest

from src.tournaments.compare_predictor_bridge import ComparePrediction
from src.tournaments.seeding_sheet import CohortSheet, SheetTeam
from src.tournaments.seeding_tiers import TierEntrant, build_cheat_sheet_analysis
from src.tournaments.seeding_workbook import build_seeding_workbook, validate_seeding_workbook, workbook_sheet_titles


def _prediction(margin: float, blowout: float = 0.1) -> ComparePrediction:
    return ComparePrediction("team_a", 0.6, 0.3, 0.1, {"teamA": 3, "teamB": 1}, margin, abs(margin), blowout)


@pytest.mark.parametrize("offset", [0.0, 1e-15, -1e-15])
def test_cheat_sheet_requires_all_available_windows_for_a_break(offset):
    ids = ["a", "b", "c", "d", "e", "f", "g", "h"]
    entrants = [TierEntrant(value, value, 1 - index * 0.05 - (.2 if index >= 4 else 0) + offset)
                for index, value in enumerate(ids)]
    predictions = {
        (first, second): _prediction(3 if ids.index(first) < 4 <= ids.index(second) else 0.1, 0.6)
        for first, second in combinations(ids, 2)
    }
    analysis = build_cheat_sheet_analysis(entrants, predictions)
    assert [item.after_seed for item in analysis.breaks] == [4]
    assert analysis.breaks[0].supported_windows == (3, 4, 5)
    assert all("division" not in note.lower() or "not" in note.lower() for note in analysis.notes)


def test_workbook_has_editable_director_columns_and_literal_names():
    cohort = CohortSheet(
        "u10", "Male",
        (SheetTeam("=Formula-like name", "Club", 0.82, 2, entrant_id="1"),),
        (),
    )
    payload = build_seeding_workbook("Event", [cohort], generated_on="2026-09-18", ranking_run="2026-09-17")
    validate_seeding_workbook(payload, ["U10 Boys"])
    workbook = load_workbook(BytesIO(payload), data_only=False)
    sheet = workbook["U10 Boys"]
    assert sheet["B7"].value == "=Formula-like name"
    assert sheet["B7"].data_type == "s"
    assert sheet.freeze_panes == "D7"
    assert sheet.auto_filter.ref is None
    assert [table.autoFilter.ref for table in sheet.tables.values()] == ["A6:K7"]
    assert [sheet.cell(6, column).value for column in range(8, 12)] == [
        "Final division", "Pool", "Final seed", "Director notes"
    ]


@pytest.mark.parametrize("event_name", ["=1+1", "+Event", "-Event", "@Event"])
def test_event_title_is_literal_text_after_reopening(event_name):
    cohort = CohortSheet("u10", "Male", (SheetTeam("Team", "Club", .8, entrant_id="0"),), ())
    payload = build_seeding_workbook(event_name, [cohort], generated_on="2026-09-19", ranking_run="2026-09-19")
    workbook = load_workbook(BytesIO(payload), data_only=False)
    assert workbook.active["A1"].value == event_name
    assert workbook.active["A1"].data_type == "s"
    assert load_workbook(BytesIO(payload), data_only=True).active["A1"].value == event_name


def test_unassigned_gender_cohort_stays_distinct_from_girls():
    sheets = [CohortSheet("u10", gender, (), (SheetTeam(name, "Club", entrant_id=str(index)),))
              for index, (gender, name) in enumerate((("Female", "Girls team"), ("", "Review team")))]
    payload = build_seeding_workbook("Draft Cup", sheets, generated_on="2026-09-19", ranking_run="unknown")
    validate_seeding_workbook(payload, workbook_sheet_titles(sheets))
    workbook = load_workbook(BytesIO(payload))
    assert workbook.sheetnames == ["U10 Girls", "U10 Unspecified gender"]
    assert workbook.worksheets[1]["A2"].value == "U10 Unspecified gender · 1 accepted teams"
    assert workbook.worksheets[1]["B7"].value == "Review team"


def test_cohort_notes_match_pdf_and_stay_outside_sortable_team_rows():
    from src.tournaments.seeding_sheet import render_sheet_html

    entrants = [TierEntrant(str(index), f"Team {index}", .8) for index in range(3)]
    analysis = build_cheat_sheet_analysis(
        entrants, {(a.entrant_id, b.entrant_id): _prediction(.1) for a, b in combinations(entrants, 2)},
    )
    cohort = CohortSheet("u10", "Female", tuple(
        SheetTeam(item.team_name, "Club", .8, entrant_id=item.entrant_id) for item in entrants
    ), (), analysis)
    notes = {("u10", "Female"): "=Director note\nKeep these words intact."}
    payload = build_seeding_workbook("Event", [cohort], generated_on="2026-09-19",
                                     ranking_run="2026-09-19", operator_notes=notes)
    validate_seeding_workbook(payload, ["U10 Girls"])
    sheet = load_workbook(BytesIO(payload)).active
    document = render_sheet_html("Event", [cohort], generated_on="2026-09-19",
                                 ranking_run="2026-09-19", operator_notes=notes)
    for note in (*analysis.notes, notes[("u10", "Female")]):
        cell = next(cell for row in sheet for cell in row if cell.value == note)
        assert cell.data_type == "s"
        assert note in document
        assert cell.row > 9
        assert str(sheet.max_row) in str(sheet.print_area)
    assert sheet.auto_filter.ref is None
    assert [table.autoFilter.ref for table in sheet.tables.values()] == ["A6:K9"]
    assert "Seeded: 3" in sheet["A5"].value
    assert all(sheet.cell(row, 11).value is None for row in range(7, 10))


def _resaved_with_second_sheet_changed(payload, change):
    workbook = load_workbook(BytesIO(payload))
    change(workbook.worksheets[1])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _set_table_filter(sheet, ref):
    for table in sheet.tables.values():
        table.autoFilter = None if ref is None else AutoFilter(ref=ref)


@pytest.mark.parametrize("change", [
    lambda sheet: setattr(sheet.auto_filter, "ref", "A1:K20"),
    lambda sheet: _set_table_filter(sheet, None),
    lambda sheet: _set_table_filter(sheet, "A6:C7"),
], ids=["sheet-filter-overlapping-the-table", "table-without-its-filter", "table-filter-over-part-of-it"])
def test_validator_rejects_a_workbook_excel_cannot_open_or_filter_on_any_sheet(change):
    cohorts = [CohortSheet(age, "Male", (SheetTeam("Team", "Club", .8, entrant_id="0"),), ()) for age in ("u10", "u11")]
    payload = build_seeding_workbook("Event", cohorts, generated_on="2026-09-18", ranking_run="2026-09-17")
    validate_seeding_workbook(payload, ["U10 Boys", "U11 Boys"])

    with pytest.raises(ValueError, match="Workbook layout is incomplete for U11 Boys"):
        validate_seeding_workbook(_resaved_with_second_sheet_changed(payload, change), ["U10 Boys", "U11 Boys"])


def test_mixed_placement_statuses_keep_identical_pdf_and_excel_order():
    import re
    from src.tournaments.seeding_sheet import render_sheet_html

    analysis = build_cheat_sheet_analysis([
        TierEntrant("seeded", "Seeded team", .9),
        TierEntrant("review", "Zebra review", .8, "Confirm the team identity."),
        TierEntrant("notfound", "Alpha not found", None, "Not found in PitchRank.",
                    review_status="Not found in PitchRank"),
        TierEntrant("unrated", "Beta unrated", None, "No current rating.", review_status="No current rating"),
    ], {})
    cohort = CohortSheet("u10", "Male", (
        SheetTeam("Seeded team", "Club", .9, entrant_id="seeded"),
        SheetTeam("Zebra review", "Club", .8, entrant_id="review"),
    ), (
        SheetTeam("Alpha not found", "Club", entrant_id="notfound"),
        SheetTeam("Beta unrated", "Club", entrant_id="unrated"),
    ), analysis)
    document = render_sheet_html("Event", [cohort], generated_on="2026-09-19", ranking_run="unknown")
    html_order = re.findall(r'data-entrant="([^"]+)"', document)
    by_id = {team.entrant_id: team.team_name for team in (*cohort.rated, *cohort.unrated)}
    payload = build_seeding_workbook("Event", [cohort], generated_on="2026-09-19", ranking_run="unknown")
    sheet = load_workbook(BytesIO(payload)).active
    assert html_order == ["seeded", "review", "notfound", "unrated"]
    assert [sheet.cell(row, 2).value for row in range(7, 11)] == [by_id[key] for key in html_order]
    assert [sheet.cell(row, 1).value for row in range(7, 11)] == [1, None, None, None]
