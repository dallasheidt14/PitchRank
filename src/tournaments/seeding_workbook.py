"""Offline Excel companion for the MatchBalance director cheat sheets."""

from __future__ import annotations

from io import BytesIO
from typing import Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.tournaments.seeding_content import DIRECTOR_LEGEND, LIMITED_HISTORY_LEGEND, build_director_cohort
from src.tournaments.seeding_pack import cohort_key, cohort_label
from src.tournaments.seeding_sheet import CohortSheet
from src.tournaments.seeding_tiers import PLACEMENT_STATUSES

FOREST = "0B5345"
FOREST_DEEP = "083E33"
YELLOW = "F4D03F"
MUTED = "5B6B66"
RULE = "D8E0DD"
BAND = "F4F7F6"


def workbook_sheet_titles(sheets: Sequence[CohortSheet]) -> list[str]:
    """Share unique cohort labels between workbook generation and validation."""
    titles = []
    used: set[str] = set()
    for sheet in sheets:
        base = cohort_label(cohort_key(sheet.age_group, sheet.gender))[:31]
        title = base
        suffix = 2
        while title.casefold() in used:
            ending = f" {suffix}"
            title = f"{base[:31 - len(ending)]}{ending}"
            suffix += 1
        titles.append(title)
        used.add(title.casefold())
    return titles


def _safe_text(value: str | None) -> str:
    """Keep imported names as literal text, including formula-like names."""
    return "" if value is None else str(value)


def _set_text(cell, value: str) -> None:
    cell.value = value
    if value[:1] in {"=", "+", "-", "@"}:
        cell.data_type = "s"
        cell._value = value


def build_seeding_workbook(
    event_name: str, sheets: Sequence[CohortSheet], *, generated_on: str, ranking_run: str,
    operator_notes: Mapping[tuple[str, str], str] | None = None,
) -> bytes:
    """Create one filterable, editable worksheet per selected cohort."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_number, (cohort, title) in enumerate(zip(sheets, workbook_sheet_titles(sheets)), 1):
        sheet = workbook.create_sheet(title)
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "D7"
        sheet.merge_cells("A1:K1")
        _set_text(sheet["A1"], _safe_text(event_name))
        sheet["A1"].font = Font(name="Arial", size=18, bold=True, color="FFFFFF")
        sheet["A1"].fill = PatternFill("solid", fgColor=FOREST)
        sheet["A1"].alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 28
        sheet.merge_cells("A2:K2")
        team_count = len(cohort.rated) + len(cohort.unrated)
        sheet["A2"] = f"{cohort_label(cohort_key(cohort.age_group, cohort.gender))} · {team_count} accepted teams"
        sheet["A2"].font = Font(bold=True, color=FOREST_DEEP)
        sheet.merge_cells("A3:K3")
        sheet["A3"] = DIRECTOR_LEGEND
        sheet["A3"].font = Font(italic=True, color=MUTED)
        sheet["A3"].alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[3].height = 30
        sheet.merge_cells("A4:K4")
        sheet["A4"] = f"Ratings as of {ranking_run} · Generated {generated_on}"
        sheet["A4"].font = Font(color=MUTED, size=9)
        content = build_director_cohort(
            cohort, (operator_notes or {}).get((cohort.age_group, cohort.gender), ""),
        )
        team_rows = content.rows
        if any(row.evidence_note for row in content.seeded):
            sheet["A3"] = DIRECTOR_LEGEND + " Limited history: " + LIMITED_HISTORY_LEGEND
            sheet.row_dimensions[3].height = 42
        sheet.merge_cells("A5:K5")
        sheet["A5"] = " · ".join(
            f"{label}: {sum(row.placement_status == label for row in team_rows)}"
            for label in PLACEMENT_STATUSES
        )
        sheet["A5"].font = Font(color=MUTED, size=9)
        headers = [
            "Suggested seed", "Team name", "Club", "PitchRank score", "State rank", "Strength marker",
            "Placement status", "Final division", "Pool", "Final seed", "Director notes",
        ]
        for column, value in enumerate(headers, 1):
            cell = sheet.cell(6, column, value)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=FOREST_DEEP)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="thin", color=YELLOW))
        for row_index, row in enumerate(team_rows, 7):
            team, marker = row.team, row.observation
            values = [
                row.seed, "\n".join(row.name_lines), _safe_text(team.club_name), row.score,
                row.state_rank, _safe_text(marker), _safe_text(row.display_status), "", "", "", "",
            ]
            for column, value in enumerate(values, 1):
                cell = sheet.cell(row_index, column, value)
                if isinstance(value, str):
                    _set_text(cell, value)
                cell.alignment = Alignment(vertical="top", wrap_text=column in {2, 3, 6, 7, 8, 9, 11})
                cell.border = Border(bottom=Side(
                    style="medium" if row.strength_break_after else "hair",
                    color=FOREST if row.strength_break_after else RULE,
                ))
                if column in {8, 9, 10, 11}:
                    cell.fill = PatternFill("solid", fgColor="FFF9DB")
                if column == 4 and value is not None:
                    cell.number_format = "0.0"
                if column in {1, 4, 5, 10}:
                    cell.alignment = Alignment(horizontal="center", vertical="top")
            if row_index % 2 == 0:
                for column in range(1, 8):
                    sheet.cell(row_index, column).fill = PatternFill("solid", fgColor=BAND)
            # Excel does not autofit wrapped rows reliably when printing.
            sheet.row_dimensions[row_index].height = 15 * max(
                sum(max(1, (len(line) + width - 1) // width) for line in str(value or "").split("\n"))
                for value, width in ((values[1], 30), (values[2], 22), (values[5], 18), (values[6], 22))
            )
        end_row = max(6, 6 + len(team_rows))
        if content.seeded:
            sheet.conditional_formatting.add(f"D7:D{6 + len(content.seeded)}", DataBarRule(
                start_type="num", start_value=0, end_type="num", end_value=100,
                color="B8D6CC", showValue=True,
            ))
        if end_row >= 7:
            table = Table(displayName=f"Cohort{sheet_number}", ref=f"A6:K{end_row}")
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium4",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=False,
                showColumnStripes=False,
            )
            sheet.add_table(table)
        notes = content.notes
        if notes:
            end_row += 2
            sheet.cell(end_row, 1, "Director notes").font = Font(bold=True, color=FOREST_DEEP)
            for note in dict.fromkeys(notes):
                end_row += 1
                sheet.merge_cells(start_row=end_row, start_column=1, end_row=end_row, end_column=11)
                cell = sheet.cell(end_row, 1)
                _set_text(cell, _safe_text(note))
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                sheet.row_dimensions[end_row].height = 15 * sum(
                    max(1, (len(line) + 139) // 140) for line in note.splitlines()
                )
        widths = [14, 32, 24, 15, 15, 20, 24, 18, 12, 14, 32]
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.print_title_rows = "1:6"
        sheet.print_area = f"A1:K{end_row}"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_margins.left = 0.25
        sheet.page_margins.right = 0.25
        sheet.page_margins.top = 0.5
        sheet.page_margins.bottom = 0.5
        sheet.oddFooter.center.text = "MatchBalance by PitchRank"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def validate_seeding_workbook(payload: bytes, expected_sheets: Sequence[str]) -> None:
    """Lightweight saved-file check used by tests and the Streamlit export."""
    workbook = load_workbook(BytesIO(payload), read_only=False, data_only=False)
    if list(workbook.sheetnames) != list(expected_sheets):
        raise ValueError("Generated workbook sheets do not match the selected cohorts.")
    for sheet in workbook.worksheets:
        tables = list(sheet.tables.values())
        # Excel reports a sheet-level filter overlapping a table's own filter as a corrupt file,
        # so filtering belongs to the table, and its filter must cover the whole table.
        if (
            sheet.freeze_panes != "D7"
            or sheet.auto_filter.ref
            or (tables and (tables[0].autoFilter is None or tables[0].autoFilter.ref != tables[0].ref))
        ):
            raise ValueError(f"Workbook layout is incomplete for {sheet.title}.")
