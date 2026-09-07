"""Build the roster-match spreadsheet from a JSON payload.

Reads the JSON described in the skill's Step 5 and writes an .xlsx with a Matches sheet
and a Notes sheet. Falls back to CSV when openpyxl is unavailable.

    python build_roster_report.py --input matches.json --out report.xlsx
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path

# (header, json key, column width)
COLUMNS = [
    ("Roster club", "roster_club", 30),
    ("Roster team", "roster_team", 44),
    ("Roster state", "roster_state", 12),
    ("PitchRank team name", "team_name", 44),
    ("PitchRank club name", "club_name", 30),
    ("State", "state_code", 7),
    ("Age group", "age_group", 10),
    ("Gender", "gender", 9),
    ("Alias provider", "providers", 24),
    ("Alias provider ID(s)", "provider_ids", 26),
    ("PitchRank UUID (team_id_master)", "team_id_master", 38),
    ("National rank", "rank_in_cohort_final", 13),
    ("State rank", "rank_in_state_final", 11),
    ("PowerScore", "power_score", 11),
    ("Ranking status", "status", 24),
    ("Ranked games (in window)", "ranked_games_window", 20),
    ("Games last 90d", "games_90d", 14),
    ("Games all time", "games_all", 14),
    ("Last played", "last_played", 12),
    ("Tier", "tier", 6),
    ("Match basis", "match_basis", 52),
    ("Alternates", "alternates", 52),
]

CENTERED = {
    "rank_in_cohort_final",
    "rank_in_state_final",
    "ranked_games_window",
    "games_90d",
    "games_all",
    "tier",
    "roster_state",
    "state_code",
    "age_group",
    "gender",
}

WRAPPED = {"match_basis", "alternates"}

HEADER_BG = "0B5345"  # PitchRank forest green
TIER_B_BG = "FEF6DC"  # prior-season row: identity inferred, not proven
TIER_C_BG = "FBE3E3"  # ambiguous: an alternate is equally good


def cell_value(value):
    """Excel takes scalars only; a list in a plural column would abort the write."""
    if isinstance(value, (list, tuple, set)):
        return "; ".join(str(v) for v in value)
    return value


def load(path: Path) -> dict:
    with path.open(encoding="utf-8") as fh:
        payload = json.load(fh)
    if not isinstance(payload, dict) or "matches" not in payload:
        raise SystemExit(f"{path}: expected a JSON object with a 'matches' key")
    return payload


def reconcile(payload: dict, roster_rows: int | None) -> None:
    """Every roster row must land in exactly one of matches or unmatched."""
    if roster_rows is None:
        return
    accounted = len(payload["matches"]) + len(payload.get("unmatched") or [])
    if accounted != roster_rows:
        raise SystemExit(
            f"roster has {roster_rows} rows but the payload accounts for {accounted} "
            f"({len(payload['matches'])} matched, "
            f"{len(payload.get('unmatched') or [])} unmatched)"
        )


def notes_lines(payload: dict) -> list[str]:
    lines = [
        payload.get("title") or "Roster match report",
        "",
        f"Season: {payload.get('season', 'unknown')}",
        f"Generated: {payload.get('generated', 'unknown')}",
        f"Last ranking run: {payload.get('ranking_run', 'unknown')}",
        "",
        "Read-only lookup. No team rows were created, edited, or merged.",
        "",
        "TIERS",
        "  A  Current-season registration row, confirmed by a fixture in the expected competition.",
        "  B  Prior-season squad row for the same club and cohort. Same players is inferred from the",
        "     name, never proven - a new season issues a new provider ID, so nothing links the rows.",
        "  C  Two or more candidates fit equally well. One is reported; the rest are in Alternates.",
        "",
        "RANKING STATUS",
        "  Active                    Ranked and published.",
        "  Not Enough Ranked Games   In the pipeline, below the games threshold to publish a rank.",
        "  Inactive                  No game inside the inactivity window, so the team is absent",
        "                            from the state board. It may still have a full slate of games.",
        "                            National and state rank are blank.",
        "",
        "COLUMNS",
        "  National rank = rank_in_cohort_final (state_rankings_view), the rank within age group",
        "                and gender nationally, frozen at the last run. rankings_full.national_rank",
        "                is always NULL - do not read it.",
        "  State rank    = rank_in_state_final, recomputed live on every read.",
        "  Ranked games (in window) = scored, non-excluded games inside the ranking window, which is",
        "                        365 days plus a 28-day grace tail (393). This is what the engine",
        "                        feeds on, and it runs well below Games all time. 12 in-window games",
        "                        is the Active threshold.",
        "  Games all time      = every game row, including unplayed future fixtures.",
    ]

    unmatched = payload.get("unmatched") or []
    if unmatched:
        lines += [
            "",
            f"NOT IN THE DATABASE ({len(unmatched)}) - omitted from the Matches sheet",
        ]
        for row in unmatched:
            club = row.get("roster_club", "")
            team = row.get("roster_team", "")
            why = row.get("note", "")
            lines.append(f"  {club} / {team}" + (f" - {why}" if why else ""))

    extra = payload.get("notes") or []
    if extra:
        lines += ["", "FINDINGS"] + [f"  {line}" for line in extra]

    return lines


def write_csv(payload: dict, out: Path) -> Path:
    out = out.with_suffix(".csv")
    with out.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow([header for header, _, _ in COLUMNS])
        for row in payload["matches"]:
            writer.writerow([cell_value(row.get(key, "")) for _, key, _ in COLUMNS])
    out.with_name(out.stem + "-notes.txt").write_text(
        "\n".join(notes_lines(payload)), encoding="utf-8"
    )
    return out


def write_xlsx(payload: dict, out: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Matches"

    header_fill = PatternFill("solid", fgColor=HEADER_BG)
    header_font = Font(bold=True, color="FFFFFF", size=11)
    tier_fills = {"B": PatternFill("solid", fgColor=TIER_B_BG), "C": PatternFill("solid", fgColor=TIER_C_BG)}

    for col, (header, _, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 30

    for r, row in enumerate(payload["matches"], start=2):
        fill = tier_fills.get(str(row.get("tier", "")).strip().upper())
        for col, (_, key, _width) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r, column=col, value=cell_value(row.get(key)))
            if fill:
                cell.fill = fill
            if key in CENTERED:
                cell.alignment = Alignment(horizontal="center")
            elif key in WRAPPED:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if key == "power_score" and row.get(key) is not None:
                cell.number_format = "0.0000"

    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(payload['matches']) + 1}"

    notes = wb.create_sheet("Notes")
    for i, line in enumerate(notes_lines(payload), start=1):
        notes.cell(row=i, column=1, value=line)
    notes.column_dimensions["A"].width = 118

    wb.save(out)
    return out


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path, help="JSON payload")
    parser.add_argument("--out", required=True, type=Path, help="Destination .xlsx")
    parser.add_argument("--roster-rows", type=int, default=None,
                        help="Row count of the source roster; asserts nothing was dropped")
    args = parser.parse_args()

    payload = load(args.input)
    reconcile(payload, args.roster_rows)
    args.out.parent.mkdir(parents=True, exist_ok=True)

    if args.out.suffix.lower() == ".csv":
        written = write_csv(payload, args.out)
    else:
        try:
            written = write_xlsx(payload, args.out)
        except Exception as exc:  # noqa: BLE001 - any write failure falls back to CSV
            written = write_csv(payload, args.out)
            print(f"xlsx write failed ({exc}) - wrote CSV instead: {written}", file=sys.stderr)

    print(written)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
